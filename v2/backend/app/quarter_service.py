import ast
import base64
import os
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from .excel_service import (
    apply_address_overrides,
    combine_workbooks,
    records_from_source,
)


@lru_cache(maxsize=1)
def _load_quarter_template_bytes():
    configured = os.getenv("QUARTER_TEMPLATE_SOURCE", "").strip()
    candidates = []
    if configured:
        candidates.append(Path(configured))
    candidates.extend([
        Path("/app/legacy_quarter_app.py"),
        Path(__file__).resolve().parents[3] / "legacy_quarter_app.py",
    ])

    source_path = next((path for path in candidates if path.exists()), None)
    if source_path is None:
        raise FileNotFoundError("기존 분기보고서 템플릿 소스를 찾을 수 없습니다.")

    source = source_path.read_text(encoding="utf-8")
    tree = ast.parse(source)
    for node in tree.body:
        if isinstance(node, (ast.Assign, ast.AnnAssign)):
            targets = node.targets if isinstance(node, ast.Assign) else [node.target]
            if any(isinstance(target, ast.Name) and target.id == "TEMPLATE_BASE64" for target in targets):
                value = ast.literal_eval(node.value)
                raw = base64.b64decode(value)
                if not raw.startswith(b"PK\x03\x04"):
                    raise RuntimeError("분기보고서 템플릿 데이터가 올바르지 않습니다.")
                return raw
    raise RuntimeError("기존 분기보고서 템플릿을 찾을 수 없습니다.")


def _aggregate_records(records):
    results = {
        (contract_type, location): [0, 0]
        for contract_type in ("공사", "용역", "물품")
        for location in (1, 2, 3)
    }
    location_map = {
        "충남도내(관내)": 1,
        "충남도내(관외)": 2,
        "타시도": 3,
    }
    for record in records:
        contract_type = record.get("목적물", "")
        location = location_map.get(record.get("소재지", ""), 3)
        if contract_type not in ("공사", "용역", "물품"):
            continue
        amount = int(record.get("계약금액", 0) or 0)
        results[(contract_type, location)][0] += 1
        results[(contract_type, location)][1] += amount
    return results


def build_quarter_report(
    file_payloads,
    target_amount,
    target_region,
    business_addresses=None,
    row_addresses=None,
):
    df, headers = combine_workbooks(file_payloads)
    apply_address_overrides(
        df,
        headers,
        business_addresses=business_addresses,
        row_addresses=row_addresses,
    )
    records = records_from_source(df, headers, target_amount, target_region)
    results = _aggregate_records(records)

    workbook = load_workbook(BytesIO(_load_quarter_template_bytes()))
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=1, max_row=4):
        for cell in row:
            if type(cell).__name__ != "MergedCell" and isinstance(cell.value, str) and "천안" in cell.value:
                cell.value = cell.value.replace("천안", target_region)

    sheet["B5"] = results[("공사", 1)][0]
    sheet["B6"] = results[("공사", 1)][1]
    sheet["C5"] = results[("공사", 2)][0]
    sheet["C6"] = results[("공사", 2)][1]
    sheet["D5"] = results[("공사", 3)][0]
    sheet["D6"] = results[("공사", 3)][1]
    sheet["F5"] = results[("용역", 1)][0]
    sheet["F6"] = results[("용역", 1)][1]
    sheet["G5"] = results[("용역", 2)][0]
    sheet["G6"] = results[("용역", 2)][1]
    sheet["H5"] = results[("용역", 3)][0]
    sheet["H6"] = results[("용역", 3)][1]
    sheet["J5"] = results[("물품", 1)][0]
    sheet["J6"] = results[("물품", 1)][1]
    sheet["K5"] = results[("물품", 2)][0]
    sheet["K6"] = results[("물품", 2)][1]
    sheet["L5"] = results[("물품", 3)][0]
    sheet["L6"] = results[("물품", 3)][1]

    output = BytesIO()
    workbook.save(output)
    return {
        "bytes": output.getvalue(),
        "record_count": len(records),
        "total_amount": sum(value[1] for value in results.values()),
        "results": results,
    }
