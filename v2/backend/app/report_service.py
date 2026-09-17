import base64
import os
import re
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .excel_service import (
    COMPETITION_METHODS,
    CONTRACT_METHODS,
    PURCHASE_PURPOSES,
    apply_address_overrides,
    categorize_region_code,
    combine_workbooks,
    extract_source_metadata,
    format_money,
    normalize_contract_type,
    records_from_source,
    region_label,
)

REVIEW_SHEET = "1-4. 기초자료(지역경제활성화)"
REQUIRED_SHEETS = [
    "1-1. 총괄(공사)",
    "1-2. 총괄(용역)",
    "1-3. 총괄(물품)",
    REVIEW_SHEET,
]
PART_SPECS = [
    ("official_01.b64", 8000),
    ("official_02.b64", 8000),
    ("official_03a.b64", 4000),
    ("official_03b.b64", 4000),
    ("official_04.b64", 4948),
]


def _template_parts_dir():
    configured = os.getenv("TEMPLATE_PARTS_DIR", "").strip()
    if configured:
        return Path(configured)

    container_path = Path("/app/template_parts")
    if container_path.exists():
        return container_path

    repo_root = Path(__file__).resolve().parents[3]
    return repo_root / "template_parts"


def _load_halfyear_template():
    part_dir = _template_parts_dir()
    chunks = []
    for name, expected_len in PART_SPECS:
        path = part_dir / name
        if not path.exists():
            raise FileNotFoundError(f"공식 반기양식 내장 데이터가 없습니다: {name}")
        text = path.read_text(encoding="utf-8").strip()
        if len(text) < expected_len:
            raise RuntimeError(f"공식 반기양식 데이터가 불완전합니다: {name}")
        chunks.append(text[:expected_len])

    try:
        raw = base64.b64decode("".join(chunks), validate=True)
        if not raw.startswith(b"PK\x03\x04"):
            raise ValueError("XLSX ZIP 헤더가 올바르지 않습니다.")
        workbook = load_workbook(BytesIO(raw))
    except Exception as exc:
        raise RuntimeError(f"공식 반기양식을 복원하지 못했습니다: {exc}") from exc

    if not set(REQUIRED_SHEETS).issubset(set(workbook.sheetnames)):
        raise RuntimeError("공식 반기양식의 4개 시트를 확인하지 못했습니다.")

    ws1 = workbook["1-1. 총괄(공사)"]
    ws3 = workbook["1-3. 총괄(물품)"]
    ws4 = workbook[REVIEW_SHEET]
    if (
        ws1["A2"].value != "순"
        or ws1["F2"].value != "총괄"
        or ws3["N2"].value != "구입목적별 "
        or "<작성방법>" not in str(ws4["A2"].value or "")
        or ws4["J5"].value != "금액단위: 원"
    ):
        raise RuntimeError("복원된 반기양식이 확정 양식 구조와 일치하지 않습니다.")

    return workbook


def _find_last_data_row(ws, start_row=7, scan_limit=5000):
    last = start_row - 1
    limit = min(ws.max_row, scan_limit)
    blank_run = 0
    for row in range(start_row, limit + 1):
        has_value = any(
            ws.cell(row, col).value not in (None, "")
            for col in range(1, min(ws.max_column, 15) + 1)
        )
        if has_value:
            last = row
            blank_run = 0
        else:
            blank_run += 1
            if last >= start_row and blank_run >= 100:
                break
    return last


def _clear_data_values(ws, end_row):
    for row in range(7, max(7, end_row) + 1):
        for col in range(1, 16):
            ws.cell(row, col).value = None


def _reset_review_validations(ws, end_row):
    ws.data_validations.dataValidation = []
    end_row = max(end_row, 1000)
    validations = [
        ("D", ["학교(유)", "학교(초)", "학교(중)", "학교(고)", "학교(특수)", "교육지원청", "직속기관"]),
        ("E", ["공사", "용역", "물품"]),
        ("F", CONTRACT_METHODS),
        ("G", COMPETITION_METHODS),
        ("M", ["충남도내(관내)", "충남도내(관외)", "타시도"]),
        ("N", PURCHASE_PURPOSES),
    ]
    for col, items in validations:
        validation = DataValidation(
            type="list",
            formula1='"' + ",".join(items) + '"',
            allow_blank=True,
        )
        ws.add_data_validation(validation)
        validation.add(f"{col}7:{col}{end_row}")


def _amount_label(target_amount):
    if target_amount >= 10000 and target_amount % 10000 == 0:
        return f"{target_amount // 10000:,}만원"
    return f"{target_amount:,}원"


def _review_title(report_year, report_label, target_amount, start_date, end_date):
    return (
        f"붙임1-4. {report_year}년 {report_label} 지역경제활성화 추진 실적"
        f"[기준(건당 {_amount_label(target_amount)} 이상), "
        f"{start_date.year}. {start_date.month}. {start_date.day}. ~ "
        f"{end_date.year}. {end_date.month}. {end_date.day}.]"
    )


def _apply_template_row_style(ws, row, template_row=7):
    if row == template_row:
        return
    ws.row_dimensions[row].height = ws.row_dimensions[template_row].height
    for col in range(1, 16):
        src = ws.cell(template_row, col)
        dst = ws.cell(row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)


def _write_review_records(ws, records, target_region, institution_name, school_level):
    for index, record in enumerate(records, start=1):
        row = index + 6
        _apply_template_row_style(ws, row)
        values = [
            index,
            target_region,
            institution_name,
            school_level,
            record.get("목적물", ""),
            record.get("계약방법", ""),
            record.get("견적경쟁방법", ""),
            record.get("계약명", ""),
            record.get("계약일자", ""),
            format_money(record.get("계약금액", 0)),
            record.get("업체명", ""),
            record.get("주소", ""),
            record.get("소재지", ""),
            record.get("구입목적", ""),
            record.get("비고", ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value
        ws.cell(row, 10).number_format = "#,##0"
        if isinstance(ws.cell(row, 9).value, datetime):
            ws.cell(row, 9).number_format = "yyyy-mm-dd"


def build_review_workbook(
    file_payloads,
    target_amount,
    target_region,
    business_addresses=None,
    row_addresses=None,
    report_year=2026,
    report_label="상반기",
    start_date=date(2026, 1, 1),
    end_date=date(2026, 7, 31),
):
    df, headers = combine_workbooks(file_payloads)
    filled_count = apply_address_overrides(
        df,
        headers,
        business_addresses=business_addresses,
        row_addresses=row_addresses,
    )
    institution_name, school_level = extract_source_metadata(df, headers)
    records = records_from_source(df, headers, target_amount, target_region)

    workbook = _load_halfyear_template()
    for name in list(workbook.sheetnames):
        if name != REVIEW_SHEET:
            workbook.remove(workbook[name])

    ws = workbook[REVIEW_SHEET]
    old_last = _find_last_data_row(ws)
    _clear_data_values(ws, max(old_last, len(records) + 6))
    ws["A1"] = _review_title(
        report_year, report_label, target_amount, start_date, end_date
    )
    ws["A4"] = (
        "◑ 자료관리목록 업로드 → 주소 보완 → 검토용 기초자료 생성 → "
        "사용자 검토 → 최종 반기보고서 생성"
    )
    ws["J4"] = "지역경제활성화 자동 집계 시스템 · Cloud Run v2"
    _write_review_records(ws, records, target_region, institution_name, school_level)
    _reset_review_validations(ws, len(records) + 6)
    ws.auto_filter.ref = f"A6:O{max(7, len(records) + 6)}"
    ws.freeze_panes = "A7"

    output = BytesIO()
    workbook.save(output)

    unresolved = sum(1 for record in records if not str(record.get("주소", "")).strip())
    return {
        "bytes": output.getvalue(),
        "record_count": len(records),
        "filled_address_count": filled_count,
        "unresolved_address_count": unresolved,
        "institution": institution_name,
        "school_level": school_level,
    }


def _find_header_row(ws):
    for row in range(1, min(ws.max_row, 30) + 1):
        values = [
            str(ws.cell(row, col).value or "").replace("\n", "").strip()
            for col in range(1, min(ws.max_column, 20) + 1)
        ]
        if "순" in values and any("계약금액" in value for value in values) and "소재지" in values:
            return row
    return None


def parse_review_workbook(review_bytes):
    try:
        workbook = load_workbook(BytesIO(review_bytes))
    except Exception as exc:
        raise ValueError(f"검토용 Excel 파일을 읽을 수 없습니다: {exc}") from exc

    source_ws = next(
        (workbook[name] for name in workbook.sheetnames if "1-4" in name or "기초자료" in name),
        workbook.active,
    )
    header_row = _find_header_row(source_ws)
    if header_row is None:
        raise ValueError("1-4 기초자료의 열 제목 행을 찾을 수 없습니다.")

    header_map = {}
    for col in range(1, source_ws.max_column + 1):
        text = str(source_ws.cell(header_row, col).value or "").replace("\n", "").strip()
        if text:
            header_map[text] = col

    def col_contains(text):
        return next((col for header, col in header_map.items() if text in header), None)

    cols = {
        "순": col_contains("순"),
        "지역": col_contains("지역"),
        "기관명": col_contains("기관명"),
        "급별": col_contains("급별"),
        "목적물": col_contains("목적물"),
        "계약방법": col_contains("계약방법"),
        "견적경쟁방법": col_contains("견적/경쟁방법") or col_contains("견적"),
        "계약명": col_contains("계약명"),
        "계약일자": col_contains("계약일자"),
        "계약금액": col_contains("계약금액"),
        "업체명": col_contains("업체명"),
        "주소": col_contains("주소"),
        "소재지": col_contains("소재지"),
        "구입목적": col_contains("구입목적"),
        "비고": col_contains("비고"),
    }

    required = ["목적물", "계약금액", "업체명", "소재지", "구입목적"]
    missing = [key for key in required if not cols[key]]
    if missing:
        raise ValueError("필수 열이 없습니다: " + ", ".join(missing))

    records = []
    blank_purpose = 0
    invalid_purpose = 0
    corrected_location = 0
    first_region = ""
    first_institution = ""
    first_level = ""
    last_data = _find_last_data_row(source_ws, header_row + 1)

    for row in range(header_row + 1, max(header_row + 1, last_data) + 1):
        ctype = normalize_contract_type(source_ws.cell(row, cols["목적물"]).value)
        amount = format_money(source_ws.cell(row, cols["계약금액"]).value)
        contract_name = (
            str(source_ws.cell(row, cols["계약명"]).value or "").strip()
            if cols["계약명"]
            else ""
        )
        company = str(source_ws.cell(row, cols["업체명"]).value or "").strip()
        if not ctype and not amount and not contract_name and not company:
            continue
        if ctype not in ["공사", "용역", "물품"]:
            continue

        region = (
            str(source_ws.cell(row, cols["지역"]).value or "").strip()
            if cols["지역"]
            else ""
        )
        institution = (
            str(source_ws.cell(row, cols["기관명"]).value or "").strip()
            if cols["기관명"]
            else ""
        )
        level = (
            str(source_ws.cell(row, cols["급별"]).value or "").strip()
            if cols["급별"]
            else ""
        )
        if not first_region and region:
            first_region = region
        if not first_institution and institution:
            first_institution = institution
        if not first_level and level:
            first_level = level

        address = (
            str(source_ws.cell(row, cols["주소"]).value or "").strip()
            if cols["주소"]
            else ""
        )
        location = str(source_ws.cell(row, cols["소재지"]).value or "").strip()
        if location not in ["충남도내(관내)", "충남도내(관외)", "타시도"]:
            location = region_label(categorize_region_code(address, region or first_region))
            corrected_location += 1

        purchase = str(source_ws.cell(row, cols["구입목적"]).value or "").strip()
        if ctype == "물품":
            if not purchase:
                purchase = "그 외"
                blank_purpose += 1
            elif purchase.startswith("도서"):
                purchase = "도서(간행물 등)"
            elif purchase not in PURCHASE_PURPOSES:
                purchase = "그 외"
                invalid_purpose += 1
        else:
            purchase = ""

        records.append({
            "순": source_ws.cell(row, cols["순"]).value if cols["순"] else len(records) + 1,
            "지역": region,
            "기관명": institution,
            "급별": level,
            "목적물": ctype,
            "계약방법": (
                str(source_ws.cell(row, cols["계약방법"]).value or "").strip()
                if cols["계약방법"]
                else ""
            ),
            "견적경쟁방법": (
                str(source_ws.cell(row, cols["견적경쟁방법"]).value or "").strip()
                if cols["견적경쟁방법"]
                else ""
            ),
            "계약명": contract_name,
            "계약일자": source_ws.cell(row, cols["계약일자"]).value if cols["계약일자"] else "",
            "계약금액": amount,
            "업체명": company,
            "주소": address,
            "소재지": location,
            "구입목적": purchase,
            "비고": (
                str(source_ws.cell(row, cols["비고"]).value or "").strip()
                if cols["비고"]
                else ""
            ),
        })

    meta = {
        "source_ws": source_ws,
        "header_row": header_row,
        "title": str(source_ws["A1"].value or ""),
        "region": first_region,
        "institution": first_institution,
        "school_level": first_level,
        "blank_purpose": blank_purpose,
        "invalid_purpose": invalid_purpose,
        "corrected_location": corrected_location,
    }
    return records, meta


def aggregate_records(records):
    base = {
        (purpose, location): [0, 0]
        for purpose in ["공사", "용역", "물품"]
        for location in [1, 2, 3]
    }
    education = {location: [0, 0] for location in [1, 2, 3]}
    books = {location: [0, 0] for location in [1, 2, 3]}
    location_map = {"충남도내(관내)": 1, "충남도내(관외)": 2, "타시도": 3}

    for record in records:
        purpose = record["목적물"]
        location = location_map.get(record["소재지"], 3)
        amount = format_money(record["계약금액"])
        base[(purpose, location)][0] += 1
        base[(purpose, location)][1] += amount

        if purpose == "물품" and record["구입목적"] == "교육용":
            education[location][0] += 1
            education[location][1] += amount
        if purpose == "물품" and str(record["구입목적"]).startswith("도서"):
            books[location][0] += 1
            books[location][1] += amount

    return base, education, books


def _copy_top_values(source_ws, target_ws, rows=6, cols=15):
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            if type(target_ws.cell(row, col)).__name__ != "MergedCell":
                target_ws.cell(row, col).value = source_ws.cell(row, col).value


def _write_final_review_sheet(ws, records, meta):
    old_last = _find_last_data_row(ws)
    _clear_data_values(ws, max(old_last, len(records) + 6))
    source_ws = meta.get("source_ws")
    if source_ws is not None:
        _copy_top_values(source_ws, ws)

    for index, record in enumerate(records, start=1):
        row = index + 6
        _apply_template_row_style(ws, row)
        values = [
            index,
            record.get("지역", ""),
            record.get("기관명", ""),
            record.get("급별", ""),
            record.get("목적물", ""),
            record.get("계약방법", ""),
            record.get("견적경쟁방법", ""),
            record.get("계약명", ""),
            record.get("계약일자", ""),
            format_money(record.get("계약금액", 0)),
            record.get("업체명", ""),
            record.get("주소", ""),
            record.get("소재지", ""),
            record.get("구입목적", ""),
            record.get("비고", ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row, col).value = value
        ws.cell(row, 10).number_format = "#,##0"
        if isinstance(ws.cell(row, 9).value, datetime):
            ws.cell(row, 9).number_format = "yyyy-mm-dd"

    _reset_review_validations(ws, len(records) + 6)
    ws.auto_filter.ref = f"A6:O{max(7, len(records) + 6)}"
    ws.freeze_panes = "A7"


def _parse_year_label(title):
    match = re.search(r"(\d{4})년\s*([^\s]+)", title or "")
    return (
        (match.group(1), match.group(2))
        if match
        else (str(datetime.now().year), "반기")
    )


def _fill_pair_group(ws, row, start_col, values):
    total_count = 0
    total_amount = 0
    col = start_col
    for location in (1, 2, 3):
        count, amount = values[location]
        ws.cell(row, col).value = count
        ws.cell(row, col + 1).value = format_money(amount)
        total_count += count
        total_amount += amount
        col += 2
    ws.cell(row, col).value = total_count
    ws.cell(row, col + 1).value = format_money(total_amount)


def _fill_summary_sheet(ws, purpose, values, region, institution, school_level, title):
    ws["A1"] = title
    ws["A5"] = 1
    ws["B5"] = region
    ws["C5"] = institution
    ws["D5"] = school_level
    ws["E5"] = purpose
    _fill_pair_group(ws, 5, 6, values)


def build_final_halfyear_report(review_bytes):
    records, meta = parse_review_workbook(review_bytes)
    if not records:
        raise ValueError("집계할 기초자료가 없습니다.")

    base, education, books = aggregate_records(records)
    workbook = _load_halfyear_template()
    year, label = _parse_year_label(meta.get("title", ""))
    region = meta.get("region", "")
    institution = meta.get("institution", "")
    school_level = meta.get("school_level", "")

    _fill_summary_sheet(
        workbook["1-1. 총괄(공사)"],
        "공사",
        {location: base[("공사", location)] for location in (1, 2, 3)},
        region,
        institution,
        school_level,
        f"붙임 1-1. {year}년 {label} 지역경제활성화 추진 실적(총괄+공사)",
    )
    _fill_summary_sheet(
        workbook["1-2. 총괄(용역)"],
        "용역",
        {location: base[("용역", location)] for location in (1, 2, 3)},
        region,
        institution,
        school_level,
        f"붙임 1-2. {year}년 {label} 지역경제활성화 추진 실적(총괄+용역)",
    )

    ws3 = workbook["1-3. 총괄(물품)"]
    ws3["A1"] = f"붙임 1-3. {year}년 {label} 지역경제활성화 추진 실적(총괄+물품)"
    ws3["A6"] = 1
    ws3["B6"] = region
    ws3["C6"] = institution
    ws3["D6"] = school_level
    ws3["E6"] = "물품"
    _fill_pair_group(
        ws3,
        6,
        6,
        {location: base[("물품", location)] for location in (1, 2, 3)},
    )
    _fill_pair_group(ws3, 6, 14, education)
    _fill_pair_group(ws3, 6, 22, books)

    _write_final_review_sheet(workbook[REVIEW_SHEET], records, meta)

    output = BytesIO()
    workbook.save(output)
    return {
        "bytes": output.getvalue(),
        "record_count": len(records),
        "region": region,
        "institution": institution,
        "school_level": school_level,
        "year": year,
        "label": label,
        "blank_purpose": meta.get("blank_purpose", 0),
        "invalid_purpose": meta.get("invalid_purpose", 0),
        "corrected_location": meta.get("corrected_location", 0),
    }
