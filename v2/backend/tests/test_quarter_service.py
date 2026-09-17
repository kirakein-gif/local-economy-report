import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from app.quarter_service import build_quarter_report


class QuarterServiceTests(unittest.TestCase):
    def _source_workbook_bytes(self, include_blank_address=False):
        headers = [f"열{i + 1}" for i in range(21)]
        headers[1] = "목적물"
        headers[2] = "계약방법"
        headers[3] = "견적/경쟁방법"
        headers[4] = "계약명"
        headers[5] = "계약일자"
        headers[6] = "계약금액(원)"
        headers[8] = "계약기관"
        headers[16] = "업체명"
        headers[18] = "사업자등록번호"
        headers[20] = "주소"

        rows = []
        source_rows = [
            ("공사", 700000, "111-11-11111", "충청남도 천안시 테스트로 1"),
            ("용역", 800000, "222-22-22222", "충청남도 아산시 테스트로 2"),
            ("물품", 900000, "333-33-33333", "서울특별시 테스트로 3"),
        ]
        if include_blank_address:
            source_rows.append(("물품", 600000, "444-44-44444", ""))

        for ctype, amount, biz, address in source_rows:
            row = [""] * 21
            row[1] = ctype
            row[2] = "수의계약"
            row[3] = "1인수의(단일견적)"
            row[4] = f"{ctype} 테스트"
            row[5] = "2026-09-01"
            row[6] = amount
            row[8] = "천안버들유치원"
            row[16] = f"{ctype}업체"
            row[18] = biz
            row[20] = address
            rows.append(row)

        output = BytesIO()
        pd.DataFrame(rows, columns=headers).to_excel(output, index=False)
        return output.getvalue()

    def test_quarter_report_uses_existing_template_and_aggregates_regions(self):
        result = build_quarter_report(
            [self._source_workbook_bytes()],
            target_amount=500000,
            target_region="천안",
        )
        self.assertEqual(result["record_count"], 3)
        self.assertEqual(result["total_amount"], 2400000)

        workbook = load_workbook(BytesIO(result["bytes"]))
        sheet = workbook.active
        self.assertEqual(sheet["B5"].value, 1)
        self.assertEqual(sheet["B6"].value, 700000)
        self.assertEqual(sheet["G5"].value, 1)
        self.assertEqual(sheet["G6"].value, 800000)
        self.assertEqual(sheet["L5"].value, 1)
        self.assertEqual(sheet["L6"].value, 900000)

    def test_quarter_report_still_downloads_with_blank_address(self):
        result = build_quarter_report(
            [self._source_workbook_bytes(include_blank_address=True)],
            target_amount=500000,
            target_region="천안",
        )
        self.assertEqual(result["record_count"], 4)
        self.assertEqual(result["total_amount"], 3000000)

        workbook = load_workbook(BytesIO(result["bytes"]))
        sheet = workbook.active
        # 주소 미확인 건은 기존 규칙대로 타시도에 임시 분류되어 보고서 생성은 계속됩니다.
        self.assertEqual(sheet["L5"].value, 2)
        self.assertEqual(sheet["L6"].value, 1500000)


if __name__ == "__main__":
    unittest.main()
