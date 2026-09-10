import unittest
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from app.report_service import (
    REQUIRED_SHEETS,
    REVIEW_SHEET,
    build_final_halfyear_report,
    build_review_workbook,
)


class ReportServiceTests(unittest.TestCase):
    def _source_workbook_bytes(self):
        headers = [f"열{i + 1}" for i in range(21)]
        headers[1] = "목적물"
        headers[2] = "계약방법"
        headers[3] = "견적/경쟁방법"
        headers[4] = "계약명"
        headers[5] = "계약일자"
        headers[6] = "계약금액(원)"
        headers[8] = "계약기관"
        headers[16] = "업체명"
        headers[18] = "사업자등록번호"
        headers[20] = "주소"

        row = [""] * 21
        row[1] = "물품"
        row[2] = "수의계약"
        row[3] = "1인수의(단일견적)"
        row[4] = "학습준비물 구입"
        row[5] = "2026-03-10"
        row[6] = 750000
        row[8] = "천안버들유치원"
        row[16] = "테스트업체"
        row[18] = "123-45-67890"
        row[20] = ""

        output = BytesIO()
        pd.DataFrame([row], columns=headers).to_excel(output, index=False)
        return output.getvalue()

    def _review_result(self):
        return build_review_workbook(
            [self._source_workbook_bytes()],
            target_amount=500000,
            target_region="천안",
            business_addresses={"1234567890": "충청남도 천안시 테스트로 1"},
            row_addresses={},
            report_year=2026,
            report_label="상반기",
            start_date=date(2026, 1, 1),
            end_date=date(2026, 7, 31),
        )

    def test_review_workbook_applies_address_override(self):
        result = self._review_result()

        self.assertEqual(result["record_count"], 1)
        self.assertEqual(result["unresolved_address_count"], 0)
        self.assertGreaterEqual(result["filled_address_count"], 1)

        workbook = load_workbook(BytesIO(result["bytes"]))
        self.assertEqual(workbook.sheetnames, [REVIEW_SHEET])
        sheet = workbook[REVIEW_SHEET]
        self.assertEqual(sheet.cell(7, 12).value, "충청남도 천안시 테스트로 1")
        self.assertEqual(sheet.cell(7, 13).value, "충남도내(관내)")
        self.assertEqual(sheet.cell(7, 14).value, "교육용")

    def test_final_report_contains_four_official_sheets(self):
        review = self._review_result()
        final = build_final_halfyear_report(review["bytes"])

        self.assertEqual(final["record_count"], 1)
        self.assertEqual(final["region"], "천안")
        self.assertEqual(final["institution"], "천안버들유치원")
        self.assertEqual(final["year"], "2026")
        self.assertEqual(final["label"], "상반기")

        workbook = load_workbook(BytesIO(final["bytes"]))
        self.assertEqual(workbook.sheetnames, REQUIRED_SHEETS)

        review_sheet = workbook[REVIEW_SHEET]
        self.assertEqual(review_sheet.cell(7, 12).value, "충청남도 천안시 테스트로 1")

        goods = workbook["1-3. 총괄(물품)"]
        self.assertEqual(goods["B6"].value, "천안")
        self.assertEqual(goods["C6"].value, "천안버들유치원")
        self.assertEqual(goods["E6"].value, "물품")
        self.assertEqual(goods["F6"].value, 1)
        self.assertEqual(goods["G6"].value, 750000)
        self.assertEqual(goods["N6"].value, 1)
        self.assertEqual(goods["O6"].value, 750000)


if __name__ == "__main__":
    unittest.main()
